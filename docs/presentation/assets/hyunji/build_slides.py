"""현지 파트 슬라이드(04-③ ×2 · 04-⑥ ×2 · 04-⑧ ×1)를 template.pptx 에서 만든다.

    python3 figures.py && python3 build_slides.py      # python-pptx · matplotlib 필요

- 템플릿 6 장(결과 제시 예시)의 틀(배경 · 머리 상자 · "04 프로젝트 수행 경과" · 부제 막대 · 오른쪽 위 라벨)만 복사하고
  내용은 새로 넣는다. 템플릿의 글꼴 · 색은 바꾸지 않는다(부제 · 라벨은 기존 글자 서식에 글만 바꾼다).
- 각 장의 노트에 발표 대사가 있다. 수치의 출처는 대사와 docs/phase1/detection-safety-geometry.md 4 · 5 장.
"""
import copy
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.package import _Relationship
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[3]
TEMPLATE = ROOT / 'docs' / 'presentation' / 'template.pptx'
OUT = ROOT / 'docs' / 'presentation' / 'slides-hyunji.pptx'
FONT = '맑은 고딕'
BLUE, RED, GRAY, DARK, LIGHT = (RGBColor(0x33, 0x78, 0xC8), RGBColor(0xC0, 0x39, 0x2B), RGBColor(0x7F, 0x7F, 0x7F),
                                RGBColor(0x26, 0x26, 0x26), RGBColor(0xF2, 0xF5, 0xFA))
KEEP = {'그래픽 98', '사각형: 둥근 한쪽 모서리 38', '그룹 2', 'TextBox 1', 'TextBox 3', '그룹 13', 'TextBox 15', '그림 34'}
FRAME_INDEX = 5          # 템플릿 6 장
LEFT, WIDTH = 0.5, 12.33


def frame_slide(prs, src, subtitle, label):
    new = prs.slides.add_slide(src.slide_layout)
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        new.part.rels._rels[rId] = _Relationship(new.part.partname.baseURI, rId, rel.reltype, RTM.INTERNAL,
                                                 rel.target_part)
    for shape in src.shapes:
        if shape.name in KEEP:
            new.shapes._spTree.insert_element_before(copy.deepcopy(shape._element), 'p:extLst')
    for shape in new.shapes:
        if shape.name == '그룹 13':
            box = next(s for s in shape.shapes if s.has_text_frame)
            set_text_keep_format(box.text_frame, subtitle)
        elif shape.name == 'TextBox 15':
            set_text_keep_format(shape.text_frame, label)
    return new


def set_text_keep_format(tf, text):
    first = tf.paragraphs[0]
    run = first.runs[0] if first.runs else first.add_run()
    run.text = text
    for extra in first.runs[1:]:
        extra._r.getparent().remove(extra._r)
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)


def text_box(slide, x, y, w, h, lines, size=13, color=DARK, bold_first=False, anchor=MSO_ANCHOR.TOP, fill=None,
             line=None, align=PP_ALIGN.LEFT):
    """lines: 문자열 또는 (문자열, {size, color, bold}) 목록."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    for i, item in enumerate(lines):
        text, opt = (item, {}) if isinstance(item, str) else item
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(opt.get('size', size))
        run.font.bold = opt.get('bold', bold_first and i == 0)
        run.font.color.rgb = opt.get('color', color)
        para.space_after = Pt(opt.get('after', 3))
    return shape


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(HERE / name), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def table(slide, x, y, w, col_w, rows, size=10, header_fill=BLUE, row_h=0.3, verdict_col=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (LIGHT if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            run = para.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = r == 0 or c == 0
            color = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else DARK
            if r > 0 and verdict_col is not None and c == verdict_col:
                color = {'달성': BLUE, '미달': RED}.get(value.split(' ')[0], GRAY)
                run.font.bold = True
            run.font.color.rgb = color
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------------------------------------------
def slide_detect(prs, src):
    s = frame_slide(prs, src, '③ 접촉 판정 — 힘 센서 없이, 로봇이 추정한 외력으로 "닿음"과 "모서리"를 가린다',
                    '* 04-③ 접촉 판정 · 안전 감시 (1/2)')
    picture(s, 'detect-dataflow.png', LEFT, 2.45, w=WIDTH)
    rows = [
        ['판정', '조건 (실기 기준)', '왜 이렇게 바꿨나 (실측)'],
        ['CONTACT\n(윗면)', '하강 중 |F − F₀| > 3 N 이 연속 3 회.\nF₀ = 최근 [t−1.0, t−0.3] s 의 이동 중 평균',
         '정지 때 잡은 F₀ 로는 윗면 38 · 44 · 78 mm 위 공중에서 거짓 접촉 → 이동 기준으로 바꾼 뒤 6 회 0 회 (9/21 실기)'],
        ['EDGE\n(모서리)', 'robot_manager 스텝 모드: 한 스텝 가고 멈춘 뒤 힘을 읽어 소실을 확정',
         '움직이는 중 z 추세선으로는 밀기 4 회 중 0 회 (9/21) → 멈춘 뒤 평균으로 바꿔 9/22 네 방향 성공'],
        ['OVER_FORCE', '원시 |F| > 30 N 한 번이면 확정. 1 차 contact_detector + 2 차 safety_monitor',
         '두 겹 감시. 30 N 은 BRD 설계 출발값(실측 아님)'],
    ]
    table(s, LEFT, 4.95, WIDTH, [1.35, 4.9, 6.08], rows, size=10.5, row_h=0.42)
    text_box(s, LEFT, 6.83, WIDTH, 0.4, [('한 줄 요약: 힘 추정값은 "그 모션 · 그 자세 근처"에서만 기준이 된다. '
                                          '판정 좌표는 조건이 처음 성립한 샘플이다.', {'bold': True, 'color': BLUE})],
             size=12)
    notes(s, '접촉 판정은 contact_detector 가 합니다. M0609 는 힘 센서가 없고, get_tool_force 는 관절 전류로 추정한 '
             '모델 값입니다. 그래서 판정 · 2 차 감시 · 정지가 모두 메인 PC 안에서, 웹을 거치지 않고 끝나게 했습니다.\n'
             'CONTACT 는 처음에 정지 상태에서 잡은 기준값을 썼는데, 9/21 실기에서 윗면 38 · 44 · 78 mm 위 공중에서 거짓 '
             '접촉이 났습니다. 추정값의 치우침이 이동 방향과 자세를 따라 흐르기 때문입니다. 기준을 최근 이동 구간 평균으로 '
             '바꾼 뒤 같은 날 6 회 하강에서 거짓 접촉은 0 회였습니다.\n'
             'EDGE 는 움직이는 중에는 힘을 믿을 수 없어서, 한 스텝 가고 멈춘 뒤 힘을 읽는 스텝 모드로 바꿨고 9/22 에 네 방향을 '
             '다 잡았습니다. 대가는 시간인데, 뒤의 KPI 장에서 말씀드리겠습니다.\n'
             '(출처: docs/phase1/detection-safety-geometry.md 3 · 5.1 · 5.2 절, PR #127 · #160)')


def slide_safety(prs, src):
    s = frame_slide(prs, src, '③ 안전 감시 — 이상이 나면 메인 PC 안에서 멈추고, 래치는 사람이 보고 푼다',
                    '* 04-③ 접촉 판정 · 안전 감시 (2/2)  [뺄 수 있음]')
    picture(s, 'safety-recovery.png', LEFT, 2.5, w=WIDTH)
    text_box(s, LEFT, 4.25, 6.05, 2.9, [
        ('실기에서 실제로 일어난 것 → 조치', {'bold': True, 'color': BLUE, 'size': 14}),
        '· 9/22 샘플 공백(약 3 s 마다 316~365 ms)으로 goal 의 약 40 % 가 SAMPLE_STALE 정지 → real 한계 300 → 500 ms 예외 '
        '조정(PR #168). 원인이 풀리면 되돌린다',
        '· 시연 중 멈춤 대응표(PR #173): 래치 · 사유 확인 → /safety/reset → 안전복귀 → 새 START. 2 회 연속이면 sim 백업 영상',
        '· 기동 직후 거짓 래치 → 기동 유예 3 s(PR #99, 04-⑧)',
    ], size=12, fill=LIGHT)
    text_box(s, 6.78, 4.25, 6.05, 2.9, [
        ('감시자도 죽는다', {'bold': True, 'color': RED, 'size': 14}),
        '· 9/21 실기 중 safety_monitor 가 두 번 죽었고, 그동안 하강 5 회 · 밀기 2 회가 2 차 감시 없이 돌았다(#103)',
        '· 원인: 같은 줄에서 로거 등급을 바꿔 불러 타이머 콜백에서 예외 → 다른 줄로 분리(PR #117, 병후 발견)',
        '· 대책: /safety/status 가 끊기면 scan_manager 가 START 를 거절(PR #136) — 감시자의 생존도 누군가 본다',
    ], size=12, fill=LIGHT)
    notes(s, '두 번째 장은 이상이 났을 때의 흐름입니다. safety_monitor 가 과대 외력 · 샘플 끊김 · 하강 제한을 확정하면 '
             '웹을 거치지 않고 /robot/stop 을 부르고, 로봇 상태로 정지 완료까지 확인합니다. 래치는 조건이 사라져도 저절로 '
             '풀리지 않고, 관제자가 원인을 보고 해제한 뒤 안전복귀와 새 START 로 이어갑니다.\n'
             '실제로 9/22 에는 샘플 공백 때문에 goal 의 약 40 % 가 멈췄고, 한계를 500 ms 로 임시 조정하면서 되돌릴 조건을 '
             '같이 적었습니다. 시연 중 멈췄을 때 사람이 할 일도 표로 만들었습니다.\n'
             '오른쪽은 감시자 자신이 죽은 경우입니다. 9/21 에 safety_monitor 가 두 번 죽는 동안 하강과 밀기 7 회가 2 차 감시 '
             '없이 돌았고, 지금은 감시자가 끊기면 스캔 시작 자체를 막습니다.\n'
             '(이 장은 목차에서 [뺄 수 있음]. 출처: 파트 문서 5.4 · 5.5 절, daily 20260923 §7)')


def slide_kpi(prs, src):
    s = frame_slide(prs, src, '⑥ 실기 결과 · KPI — 치수 정확도는 달성, 탐색 시간 · 검출 하중 · 반영 지연은 미달',
                    '* 04-⑥ 실기 결과 · KPI (1/2)')
    rows = [
        ['항목', 'KPI (BRD 9 장)', '실측값 (날짜 · 조건 · 출처)', '판정', '미달 원인', '개선안'],
        ['모서리 좌표(치수) 오차', '±3 mm', '가로 +1.50 · 세로 −0.36 mm — 9/23 실기 1 회, 캘리퍼 대비(scan …3702, #180)',
         '달성', '—', '10 회 반복(TR-02)으로 재확인'],
        ['윗면 높이 오차', '±3 mm', '높이 +0.26 mm — 같은 scan', '달성', '—', '같음'],
        ['접촉 검출 하중', '5 N 이하', '임계 3 N: 평균 4.49 N, 10 회 중 2 회 초과(5.15 · 5.29)\n임계 2 N: 평균 2.76 N, 0/7 — '
         '9/23 실기(TR-01, PR #181)', '미달', '회차 기준 2/10 초과. 평균만 보면 통과처럼 보인다',
         '임계 2 N 검토 — 공중 거짓 접촉(#182)과 같이'],
        ['작업 중지 반응 시간', '1 s 이내', '706 · 600 ms — 9/22 sim, 클릭 → STOPPED 수신(TR-05)', '달성 (sim)', '—',
         '실기 측정'],
        ['관제 화면 반영 지연', '200 ms 이내', '최대 312 ms — 9/22(TR-05)', '미달', '최대값을 만든 토픽을 못 가렸다',
         '토픽별 지연 기록 재시험'],
        ['전체 탐색 시간', '120 s 이내', '438 s — 9/23 실기, 스텝 긁기 331 s · 재접근 내림 100 s', '미달',
         '정확도를 위해 멈춰서 읽는 스텝 모드', '스텝 크기 · 속도 조정(#180)'],
        ['실기 연속 성공', '(목차 항목)', '웹 START → 형상 → 웹 표시 종단 성공 — 9/23(PR #179). 연속 회차 기록은 없다', '미확인', '—',
         '리허설에서 연속 회차 기록'],
        ['반복성', '10 회 σ 0.3 mm', '미실시 (같은 점 10 회 z 산포 0.040 mm 는 다른 양)', '미실시', '—', 'TR-02'],
        ['접촉 오검출', '20 회 중 1 회 이하', '미실시 (참고: 9/23 하강 공중 거짓 접촉 3/10, #182)', '미실시', '—', '#182 수정 뒤 20 회'],
        ['품종 교체 셋업 시간', '3 분 이내', '미실시', '미실시', '—', 'TR-04'],
    ]
    table(s, LEFT, 2.42, WIDTH, [1.55, 1.2, 4.45, 0.9, 2.2, 2.03], rows, size=9.5, row_h=0.43, verdict_col=3)
    notes(s, 'KPI 표입니다. 미달을 숨기지 않고 그대로 적었습니다.\n'
             '치수 정확도는 9/23 실기에서 캘리퍼 대비 가로 +1.50, 세로 −0.36, 높이 +0.26 mm 로 ±3 mm 안입니다. 다만 한 번 '
             '잰 값이고 10 회 반복성은 아직 못 했습니다.\n'
             '검출 하중은 평균 4.49 N 이라 평균으로는 통과지만, 10 회 중 2 회가 5 N 을 넘어 회차 기준으로 미달로 적었습니다. '
             '임계를 2 N 으로 내리면 0/7 이지만 공중 거짓 접촉 문제와 같이 풀어야 합니다.\n'
             '중지 반응은 sim 에서만 쟀고, 반영 지연은 최대 312 ms 로 미달입니다. 탐색 시간은 438 s 로 목표의 3.7 배입니다. '
             '실기 연속 성공 횟수는 기록이 없어 미확인으로 두었습니다.\n'
             '(출처: TR-01_20260923(PR #181) · TR-05_20260922 · #180 · PR #179. 확인 필요: 연속 성공 회차 기록이 있으면 병후에게)')


def slide_kpi_story(prs, src):
    s = frame_slide(prs, src, '⑥ 미달을 어떻게 읽나 — 정확도를 먼저 잡고, 시간을 내줬다',
                    '* 04-⑥ 실기 결과 · KPI (2/2)')
    picture(s, 'kpi-scan-time.png', LEFT, 2.42, w=7.2)
    picture(s, 'kpi-detect-load.png', LEFT, 4.08, w=7.2)
    steps = [
        ('① 처음', 'KPI 를 "한 번 성공하면 된다"로 읽었다. 회차 기준 · 시간 기준을 늦게 챙겼다'),
        ('② 판단', '9/21 움직이는 중의 힘으로는 모서리 0/4 → 멈춰서 읽는 스텝 모드로 정확도 우선(9/22)'),
        ('③ 결과', '치수 ±3 mm 달성. 대신 탐색 438 s — 시간의 76 % 가 스텝 긁기'),
        ('④ 정직하게', '검출 하중 평균 4.49 N 은 통과처럼 보이지만 회차 2/10 초과 → 미달로 적었다'),
        ('⑤ 개선안', '스텝 크기 · 속도 조정으로 긁기 시간 줄이기(#180) · 임계 2 N + 공중 거짓 접촉 수정(#182) · '
                   '반영 지연 토픽별 재시험'),
    ]
    y = 2.42
    for head, body in steps:
        text_box(s, 7.9, y, 4.93, 0.9, [(head, {'bold': True, 'color': BLUE, 'size': 12}), (body, {'size': 11})],
                 fill=LIGHT)
        y += 0.97
    notes(s, '왜 미달이 났는지 순서대로 말씀드리겠습니다. 처음에는 KPI 를 한 번 성공하는 것으로 읽었고, 회차 기준과 시간 '
             '기준을 늦게 챙겼습니다. 9/21 에 움직이는 중의 힘으로는 모서리를 한 번도 못 잡아서, 멈춰서 힘을 읽는 스텝 '
             '모드로 바꾸며 정확도를 먼저 잡았습니다. 그 결과 치수는 목표 안에 들어왔지만 탐색 시간은 438 s 가 되었고, '
             '위 그래프처럼 그중 76 % 가 스텝 긁기입니다.\n'
             '아래 그래프는 검출 하중입니다. 빨간 막대 두 개가 5 N 을 넘은 회차이고, 이것 때문에 평균이 통과여도 미달로 '
             '적었습니다. 오른쪽은 임계 2 N 의 결과입니다.\n'
             '개선안은 스텝 크기와 속도를 조정해 긁는 시간을 줄이는 것, 임계를 내리되 공중 거짓 접촉을 같이 고치는 것, '
             '반영 지연을 토픽별로 다시 재는 것입니다.')


def slide_problem(prs, src):
    s = frame_slide(prs, src, '⑧ 기동 직후의 거짓 래치를 "기동 유예 3 s"로 없앴다 — 과대 외력은 유예하지 않는다',
                    '* 04-⑧ 문제와 해결 (현지)')
    boxes = [
        ('증상', ['9/21 아침 sim 종단: /scan/run 이 SAFETY_LATCHED(103) 로 거절', '로그 "SAMPLE_STALE: 508 ms 동안 수신 없음"',
                 '로봇은 서 있었다']),
        ('원인', ['노드 5 개가 동시에 뜨며 샘플이 508 ms 끊김', 'robot_manager 는 "위치를 모르면 이동 중"(PR #72)',
                 '→ "움직이는 중 끊김 = 정지" 규칙과 겹침. 각자 옳은 기본값 두 개의 충돌']),
        ('조치 (PR #99)', ['startup_grace_s 3.0 s 동안 최신성으로는 정지 · 래치하지 않는다',
                          '과대 외력 · 하강 제한은 유예 없이 즉시', '시험으로 고정(유예 시험 2 개 추가)']),
        ('결과', ['고치기 전: level 2 · latched · 403', '고친 뒤: level 0 · latched false (Virtual + sim 재현)',
                 'safety_monitor 시험 52 개 통과']),
    ]
    x = LEFT
    w = (WIDTH - 0.3) / 4
    for i, (head, lines) in enumerate(boxes):
        text_box(s, x, 2.42, w, 1.95, [(f'{i + 1}. {head}', {'bold': True, 'size': 14, 'color': BLUE if i else RED})]
                 + [(f'· {t}', {'size': 11}) for t in lines], fill=LIGHT)
        x += w + 0.1
    picture(s, 'problem-startup-latch.png', 1.2, 4.5, h=2.35)
    text_box(s, LEFT, 6.9, WIDTH, 0.38, [('교훈: 노드마다 옳은 보수적 기본값도 겹치면 오탐이 된다 — 노드 사이의 기본값은 같이 본다',
                                          {'bold': True, 'color': BLUE})], size=12)
    notes(s, '제가 고친 문제 하나입니다. 9/21 아침 sim 종단에서 스캔 시작이 안전 래치로 거절됐습니다. 로봇은 서 있었는데 '
             'safety_monitor 가 샘플이 508 ms 끊겼다며 래치를 걸었습니다.\n'
             '원인은 노드 다섯 개가 동시에 뜨면서 샘플이 잠깐 끊긴 것과, robot_manager 가 위치를 모르면 움직이는 중으로 '
             '보고하는 규칙이 겹친 것이었습니다. 둘 다 각자는 옳은 보수적 기본값입니다.\n'
             '그래서 기동 뒤 3 초 동안은 샘플 최신성으로는 멈추지 않게 했고, 과대 외력과 하강 제한은 유예하지 않도록 시험으로 '
             '고정했습니다. 고친 뒤에는 기동 직후 래치가 걸리지 않습니다.\n'
             '(출처: PR #99, 파트 문서 5.3 절. 그림은 개념도 — 공백의 시작 시각은 기록이 없다)')


def main():
    prs = Presentation(str(TEMPLATE))
    src = prs.slides[FRAME_INDEX]
    template_ids = list(prs.slides._sldIdLst)
    for build in (slide_detect, slide_safety, slide_kpi, slide_kpi_story, slide_problem):
        build(prs, src)
    for sld in template_ids:                       # 양식 장(표지 · 목차 · 예시)은 지운다. 표지 · 목차는 병후 파일에 있다
        prs.part.drop_rel(sld.rId)
        prs.slides._sldIdLst.remove(sld)
    prs.save(str(OUT))
    print('saved', OUT.relative_to(ROOT), len(prs.slides), 'slides')


if __name__ == '__main__':
    main()
